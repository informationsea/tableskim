#!/usr/bin/env python3
"""Generate demo data for testing TableSkim's table viewers.

Creates the following datasets in the output directory, each as CSV, .xlsx
and .parquet, plus demo.sqlite3 and demo.duckdb containing one table per
dataset:

  normal     : A well-formed table with a header and mixed value types
  ragged     : A table whose column count varies from row to row
  large      : A 10,000-row table
  dates      : Dates, date-times, times of day and durations
  多言語サンプル : Greetings and native digits in ~25 languages and scripts
               (CJK, Hangul, Cyrillic, Greek, Arabic, Hebrew, Indic, Thai, ...)
               plus Unicode edge cases (emoji sequences, combining marks,
               bidirectional text, astral-plane CJK, full/half width forms).
               The file name itself is Japanese; the .xlsx also has one sheet
               per script with a non-ASCII sheet name.

dates.xlsx additionally has a sheet showing the same instants with many
different Excel number formats (built-in, custom, CJK, elapsed time) next to
numeric formats that must not be mistaken for dates, and dates1904.xlsx holds
the same values in the 1904 date system.

normal.csv is also written compressed as normal.csv.gz, .bz2, .xz and .zst,
and large.csv as large.tsv.gz (tab-separated).

normal.csv is also written in other text encodings: normal-utf8-bom.csv,
normal-cp932.csv (and .csv.gz), normal-utf16le-bom.tsv (like Excel's
"Unicode Text"), and normal-utf16le.csv / normal-utf16be.csv without a BOM.
cp932-sample.csv holds characters that are tricky in CP932: trail bytes
that are '\\' or '|', NEC and IBM extensions, half-width katakana, and
quoted fields with commas and line breaks.

super-large.csv.xz, super-large.xlsx and super-large.parquet (zstd) hold
1,000,000 rows shaped like "large", streamed to disk, for testing
performance and memory use. They are the slowest part to generate; skip them
with --no-super-large.

normal is also written as JSON: normal.json (an array of objects),
normal.ndjson (one object per line) and normal.jsonl.gz. nested.json holds
nested objects and arrays, keys missing from some objects, and a key whose
values change type after the first rows.

Excel-only files:

  multisheet.xlsx : Several worksheets (including an empty one, a hidden one,
                    one whose data starts away from A1 and a chart sheet)
  merged.xlsx     : Merged cells in the header, horizontally, vertically and
                    as blocks, including merges reaching past the last value

Usage:
  python3 create-demo-data.py [-o OUTPUT_DIR]

Requires openpyxl, pyarrow, duckdb and zstandard (pip install -r requirements.txt).
"""

import argparse
import bz2
import csv
import datetime
import gzip
import io
import json
import lzma
import random
import sqlite3
from pathlib import Path

import duckdb
import pyarrow as pa
import pyarrow.parquet as pq
import zstandard
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font
from openpyxl.utils.datetime import CALENDAR_MAC_1904

SEED = 20261005
SUPER_LARGE_SEED = SEED + 4

FIRST_NAMES = ["Taro", "Hanako", "Ichiro", "Yuki", "Sakura", "Kenji", "Aoi",
               "Alice", "Bob", "Carol", "David", "Eve", "Frank", "Grace"]
LAST_NAMES = ["Sato", "Suzuki", "Takahashi", "Tanaka", "Watanabe", "Ito",
              "Smith", "Johnson", "Brown", "Miller", "Davis", "Wilson"]
CITIES = ["Sendai", "Tokyo", "Osaka", "Sapporo", "Fukuoka", "Nagoya",
          "New York", "London", "Paris", "Berlin"]
DEPARTMENTS = ["Engineering", "Sales", "Marketing", "Research", "Support",
               "人事部", "経理部"]


def write_csv(path: Path, rows):
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerows(rows)


def csv_bytes(rows, delimiter: str = ",") -> bytes:
    buffer = io.StringIO(newline="")
    csv.writer(buffer, delimiter=delimiter).writerows(rows)
    return buffer.getvalue().encode("utf-8")


COMPRESSORS = {
    # mtime=0 keeps the output reproducible.
    ".gz": lambda data: gzip.compress(data, mtime=0),
    ".bz2": bz2.compress,
    ".xz": lzma.compress,
    ".zst": lambda data: zstandard.ZstdCompressor().compress(data),
}


def write_compressed(path: Path, data: bytes):
    path.write_bytes(COMPRESSORS[path.suffix](data))


ENCODED_VARIANTS = [
    # file name, delimiter, encoding (Python codec), byte order mark
    ("normal-utf8-bom.csv", ",", "utf-8", b"\xef\xbb\xbf"),
    ("normal-cp932.csv", ",", "cp932", b""),
    ("normal-utf16le-bom.tsv", "\t", "utf-16-le", b"\xff\xfe"),
    ("normal-utf16le.csv", ",", "utf-16-le", b""),
    ("normal-utf16be.csv", ",", "utf-16-be", b""),
]

CP932_SAMPLE = [
    ["ID", "Text", "Note"],
    [1, "表示,ソフト", "trail byte 0x5C (\\) and a quoted comma"],
    [2, "ポイント|予定", "trail byte 0x7C (|)"],
    [3, "①②③ ㈱ №", "NEC special characters"],
    [4, "髙橋 﨑 纊", "IBM extensions"],
    [5, "ｶﾀｶﾅ ﾃｽﾄ", "half-width katakana"],
    [6, "～－￠￡￢∥", "CP932-specific mappings"],
    [7, "改行を\n含む", "quoted line break"],
]


def json_value(v):
    if isinstance(v, (datetime.date, datetime.datetime, datetime.time)):
        return v.isoformat()
    return v


def json_records(rows):
    header, *body = rows
    return [{key: json_value(v) for key, v in zip(header, row)} for row in body]


def write_json(output: Path, rows):
    records = json_records(rows)
    (output / "normal.json").write_text(json.dumps(records, ensure_ascii=False, indent=2) + "\n",
                                        encoding="utf-8")
    lines = "".join(json.dumps(r, ensure_ascii=False) + "\n" for r in records).encode("utf-8")
    (output / "normal.ndjson").write_bytes(lines)
    write_compressed(output / "normal.jsonl.gz", lines)
    nested = []
    for i in range(1, 31):
        record = {
            "id": i,
            "user": {"name": f"user{i}", "tags": ["a", "b"][: i % 3]},
            "scores": [i, i * 2, i * 3],
            # A number at first, text from row 25 on.
            "code": i if i < 25 else f"C-{i}",
        }
        if i % 4 == 0:
            record["note"] = f"note {i}"
        nested.append(record)
    (output / "nested.json").write_text(json.dumps(nested, ensure_ascii=False, indent=2) + "\n",
                                        encoding="utf-8")


def write_encoded(output: Path, rows):
    for name, delimiter, encoding, bom in ENCODED_VARIANTS:
        text = csv_bytes(rows, delimiter).decode("utf-8")
        (output / name).write_bytes(bom + text.encode(encoding))
    cp932 = csv_bytes(rows).decode("utf-8").encode("cp932")
    write_compressed(output / "normal-cp932.csv.gz", cp932)
    sample = csv_bytes(CP932_SAMPLE).decode("utf-8").encode("cp932")
    (output / "cp932-sample.csv").write_bytes(sample)


def fill_sheet(ws, rows, header: bool = True):
    for row in rows:
        ws.append(row)
    if header and rows:
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"


def write_xlsx(path: Path, rows, sheet_title: str, header: bool = True):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    fill_sheet(ws, rows, header)
    wb.save(path)


def to_rectangular(rows, header: bool):
    """Split off column names and pad ragged rows with None."""
    width = max(len(r) for r in rows)
    if header:
        columns, body = list(rows[0]), rows[1:]
    else:
        columns, body = [f"col{c}" for c in range(1, width + 1)], rows
    return columns, [list(r) + [None] * (width - len(r)) for r in body]


def to_arrow(rows, header: bool) -> pa.Table:
    columns, body = to_rectangular(rows, header)
    return pa.table({col: pa.array([r[c] for r in body])
                     for c, col in enumerate(columns)})


def write_parquet(path: Path, rows, header: bool):
    pq.write_table(to_arrow(rows, header), path)


def write_duckdb(path: Path, datasets):
    """Write each dataset as a table with native types (DATE, TIMESTAMP, ...)."""
    path.unlink(missing_ok=True)
    con = duckdb.connect(str(path))
    for name, rows, header in datasets:
        arrow_table = to_arrow(rows, header)  # noqa: F841 (referenced by DuckDB)
        con.execute(f'CREATE TABLE "{name}" AS SELECT * FROM arrow_table')
    con.execute('CREATE INDEX "large_category" ON "large" ("Category")')
    con.execute(LARGE_SUMMARY_VIEW)
    con.close()


LARGE_SUMMARY_VIEW = ('CREATE VIEW "large_summary" AS '
                      'SELECT "Category", COUNT(*) AS n, AVG("Score") AS avg_score '
                      'FROM "large" GROUP BY "Category"')


def write_sqlite(path: Path, datasets):
    """Write each dataset as a table. Ragged rows are padded with NULL."""
    path.unlink(missing_ok=True)
    con = sqlite3.connect(path)
    for name, rows, header in datasets:
        columns, body = to_rectangular(rows, header)
        width = len(columns)
        types = [sqlite_type([r[c] for r in body if c < len(r)]) for c in range(width)]
        col_defs = ", ".join(f'"{col}" {t}' for col, t in zip(columns, types))
        con.execute(f'CREATE TABLE "{name}" ({col_defs})')
        placeholders = ", ".join("?" * width)
        con.executemany(
            f'INSERT INTO "{name}" VALUES ({placeholders})',
            ([sqlite_value(v) for v in r] for r in body))
    # A view and an index so viewers can be tested against non-table objects.
    con.execute('CREATE INDEX "large_category" ON "large" ("Category")')
    con.execute(LARGE_SUMMARY_VIEW)
    con.commit()
    con.close()


def sqlite_type(values):
    values = [v for v in values if v is not None and v != ""]
    if values and all(isinstance(v, (bool, int)) for v in values):
        return "INTEGER"
    if values and all(isinstance(v, (int, float)) for v in values):
        return "REAL"
    return "TEXT"


def sqlite_value(v):
    if isinstance(v, datetime.datetime):
        return v.isoformat(sep=" ")
    if isinstance(v, (datetime.date, datetime.time)):
        return v.isoformat()
    if isinstance(v, datetime.timedelta):
        return format_duration(v)
    if isinstance(v, bool):
        return int(v)
    return v


def format_duration(v: datetime.timedelta) -> str:
    seconds = int(v.total_seconds())
    return "%d:%02d:%02d" % (seconds // 3600, seconds // 60 % 60, seconds % 60)


def random_person(rng: random.Random, i: int):
    first = rng.choice(FIRST_NAMES)
    last = rng.choice(LAST_NAMES)
    birthday = datetime.date(1960, 1, 1) + datetime.timedelta(days=rng.randrange(0, 365 * 45))
    return [
        i,
        f"{first} {last}",
        f"{first.lower()}.{last.lower()}{i}@example.com",
        rng.randint(20, 65),
        rng.choice(CITIES),
        rng.choice(DEPARTMENTS),
        round(rng.uniform(250000, 1500000), 2),
        birthday,
        rng.choice([True, False]),
    ]


PERSON_HEADER = ["ID", "Name", "Email", "Age", "City", "Department",
                 "Salary", "Birthday", "Active"]


def csv_value(v):
    if isinstance(v, (datetime.date, datetime.datetime, datetime.time)):
        return v.isoformat()
    if isinstance(v, datetime.timedelta):
        return format_duration(v)
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    return v


def to_csv_rows(rows):
    return [[csv_value(v) for v in row] for row in rows]


def make_normal(rng: random.Random):
    rows = [PERSON_HEADER]
    rows += [random_person(rng, i) for i in range(1, 51)]
    # Tricky-but-valid values: quoting, embedded commas/newlines, empty cells.
    rows += [
        [51, "Smith, John", "john.smith@example.com", 40, "Tokyo",
         'Says "hello"', 500000.0, datetime.date(1985, 4, 1), True],
        [52, "複数行\nの名前", "multi.line@example.com", 33, "仙台",
         "研究開発部", 720000.5, datetime.date(1992, 12, 31), False],
        [53, "", "", None, "", "", None, None, None],
        [54, "Leading and trailing spaces  ", "  pad@example.com", 29, " Osaka ",
         "Sales", 0, datetime.date(2000, 2, 29), True],
    ]
    return rows


def make_ragged(rng: random.Random):
    rows = [["A", "B", "C", "D", "E"]]
    lengths = [5, 3, 7, 1, 0, 5, 12, 2, 5, 4, 9, 6, 1, 15, 5, 3, 0, 8, 5, 20]
    for r, n in enumerate(lengths, start=1):
        rows.append([f"r{r}c{c}" if rng.random() > 0.1 else "" for c in range(1, n + 1)])
    # A few rows with trailing empty cells, and a very wide row at the end.
    rows.append(["trailing", "empties", "", "", ""])
    rows.append([f"w{c}" for c in range(1, 51)])
    return rows


LARGE_HEADER = ["RowNo", "UUID", "Category", "Value", "Score", "Timestamp", "Comment"]
SUPER_LARGE_ROWS = 1_000_000


def iter_large_rows(rng: random.Random, n_rows: int):
    """Data rows (without the header) of the "large" layout, one at a time."""
    categories = ["alpha", "beta", "gamma", "delta", "epsilon"]
    base = datetime.datetime(2026, 1, 1)
    for i in range(1, n_rows + 1):
        yield [
            i,
            "%08x-%04x-%04x-%04x-%012x" % (
                rng.getrandbits(32), rng.getrandbits(16), rng.getrandbits(16),
                rng.getrandbits(16), rng.getrandbits(48)),
            rng.choice(categories),
            rng.randint(-100000, 100000),
            round(rng.gauss(50, 15), 4),
            base + datetime.timedelta(seconds=rng.randrange(0, 365 * 24 * 3600)),
            " ".join(rng.choice(["lorem", "ipsum", "dolor", "sit", "amet",
                                 "consectetur", "adipiscing", "elit"])
                     for _ in range(rng.randint(1, 8))),
        ]


def make_large(rng: random.Random, n_rows: int = 10000):
    return [LARGE_HEADER] + list(iter_large_rows(rng, n_rows))


LARGE_SCHEMA = pa.schema([
    ("RowNo", pa.int64()), ("UUID", pa.string()), ("Category", pa.string()),
    ("Value", pa.int64()), ("Score", pa.float64()), ("Timestamp", pa.timestamp("us")),
    ("Comment", pa.string()),
])


def write_super_large_parquet(output: Path, seed: int, n_rows: int = SUPER_LARGE_ROWS,
                              batch_rows: int = 100_000):
    rows = iter_large_rows(random.Random(seed), n_rows)
    with pq.ParquetWriter(output / "super-large.parquet", LARGE_SCHEMA, compression="zstd") as writer:
        while batch := [row for _, row in zip(range(batch_rows), rows)]:
            columns = list(zip(*batch))
            writer.write_table(pa.table(
                {field.name: pa.array(col, field.type) for field, col in zip(LARGE_SCHEMA, columns)},
                schema=LARGE_SCHEMA))


def write_super_large(output: Path, seed: int, n_rows: int = SUPER_LARGE_ROWS):
    """Write super-large.csv.xz, .xlsx and .parquet without keeping rows in memory.

    Every file is generated from the same seed instead of buffering the rows.
    """
    with lzma.open(output / "super-large.csv.xz", "wt", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(LARGE_HEADER)
        for row in iter_large_rows(random.Random(seed), n_rows):
            writer.writerow([csv_value(v) for v in row])

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("super-large")
    ws.freeze_panes = "A2"
    header = []
    for name in LARGE_HEADER:
        cell = WriteOnlyCell(ws, value=name)
        cell.font = Font(bold=True)
        header.append(cell)
    ws.append(header)
    for row in iter_large_rows(random.Random(seed), n_rows):
        ws.append(row)
    wb.save(output / "super-large.xlsx")

    write_super_large_parquet(output, seed, n_rows)


def make_dates(rng: random.Random, n_rows: int = 60):
    """Typed date/time columns, with a few edge cases first."""
    header = ["ID", "Date", "DateTime", "Time", "Duration", "DateText"]
    rows = [header]
    edge_cases = [
        # Excel's 1900 system has a fictitious 1900-02-29; dates around it and
        # the start of each epoch are easy to get wrong.
        datetime.datetime(1900, 1, 1, 0, 0, 0),
        datetime.datetime(1900, 3, 1, 12, 0, 0),
        datetime.datetime(1904, 1, 1, 0, 0, 0),
        datetime.datetime(1970, 1, 1, 0, 0, 0),
        datetime.datetime(2000, 2, 29, 23, 59, 59),
        datetime.datetime(2026, 10, 5, 9, 30, 15, 250000),
        datetime.datetime(9999, 12, 31, 23, 59, 59),
    ]
    stamps = edge_cases + [
        datetime.datetime(1950, 1, 1) + datetime.timedelta(seconds=rng.randrange(0, 100 * 365 * 24 * 3600))
        for _ in range(n_rows - len(edge_cases))
    ]
    for i, stamp in enumerate(stamps, start=1):
        rows.append([
            i,
            stamp.date(),
            stamp,
            stamp.time().replace(microsecond=0),
            datetime.timedelta(seconds=rng.randrange(0, 200 * 3600)),
            stamp.strftime("%Y/%m/%d"),  # looks like a date, but is text
        ])
    # Missing values in every column.
    rows.append([len(stamps) + 1, None, None, None, None, ""])
    return rows


# (label, value, number format). Format None keeps openpyxl's default.
DATE_FORMAT_SAMPLES = [
    ("builtin 14 (m/d/yy)", datetime.date(2026, 10, 5), "mm-dd-yy"),
    ("builtin 22 (m/d/yy h:mm)", datetime.datetime(2026, 10, 5, 9, 30), "m/d/yy h:mm"),
    ("builtin 18 (h:mm AM/PM)", datetime.time(21, 5), "h:mm AM/PM"),
    ("builtin 21 (h:mm:ss)", datetime.time(9, 5, 30), "h:mm:ss"),
    ("builtin 46 ([h]:mm:ss)", datetime.timedelta(hours=50, minutes=3, seconds=4), "[h]:mm:ss"),
    ("ISO date", datetime.date(2026, 10, 5), "yyyy-mm-dd"),
    ("ISO date-time", datetime.datetime(2026, 10, 5, 9, 30, 15), "yyyy-mm-dd hh:mm:ss"),
    ("slash date", datetime.date(2026, 10, 5), "yyyy/mm/dd"),
    ("month name", datetime.date(2026, 10, 5), "mmm d, yyyy"),
    ("day name", datetime.date(2026, 10, 5), "dddd, mmmm dd, yyyy"),
    ("Japanese date", datetime.date(2026, 10, 5), 'yyyy"年"m"月"d"日"'),
    ("Japanese era", datetime.date(2026, 10, 5), '[$-ja-JP]ggge"年"m"月"d"日"'),
    ("locale prefix", datetime.date(2026, 10, 5), "[$-409]d-mmm-yy;@"),
    ("minutes only", datetime.time(0, 45), "mm:ss"),
    ("elapsed minutes", datetime.timedelta(minutes=125), "[mm]:ss"),
    ("fractional seconds", datetime.datetime(2026, 10, 5, 9, 30, 15, 500000), "hh:mm:ss.000"),
    ("date with color", datetime.date(2026, 10, 5), "[Red]yyyy-mm-dd"),
    ("default date", datetime.date(2026, 10, 5), None),
    ("default date-time", datetime.datetime(2026, 10, 5, 9, 30, 15), None),
    ("default time", datetime.time(9, 30, 15), None),
    # Numeric formats that must not be read as dates.
    ("general number", 45935, "General"),
    ("thousands", 1234567.891, "#,##0.00"),
    ("percent", 0.125, "0.00%"),
    ("scientific", 12345.678, "0.00E+00"),
    ("currency", -1234.5, '"¥"#,##0;[Red]"-¥"#,##0'),
    ("quoted letters", 42, '0 "days"'),
    ("escaped letters", 42, r"0\h"),
    ("text format", "2026-10-05", "@"),
]


def add_format_samples_sheet(wb: Workbook):
    ws = wb.create_sheet("formats")
    ws.append(["Label", "Value", "NumberFormat", "Kind", "ISO"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for label, value, number_format in DATE_FORMAT_SAMPLES:
        if isinstance(value, datetime.datetime):
            kind, iso = "datetime", value.isoformat(sep=" ")
        elif isinstance(value, datetime.date):
            kind, iso = "date", value.isoformat()
        elif isinstance(value, datetime.time):
            kind, iso = "time", value.isoformat()
        elif isinstance(value, datetime.timedelta):
            kind, iso = "duration", format_duration(value)
        else:
            kind, iso = ("text" if isinstance(value, str) else "number"), str(value)
        ws.append([label, value, number_format or "(default)", kind, iso])
        if number_format:
            ws.cell(row=ws.max_row, column=2).number_format = number_format
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 34


def write_dates_xlsx(path: Path, rows, epoch=None):
    wb = Workbook()
    if epoch is not None:
        wb.epoch = epoch
    ws = wb.active
    ws.title = "dates"
    # 1904-based workbooks cannot store dates before 1904.
    min_date = datetime.date(1904, 1, 2) if epoch is not None else datetime.date.min
    fill_sheet(ws, [rows[0]] + [r for r in rows[1:] if r[1] is None or r[1] >= min_date])
    add_format_samples_sheet(wb)
    wb.save(path)


def write_multisheet_xlsx(path: Path, rng: random.Random):
    wb = Workbook()

    people = wb.active
    people.title = "People"
    fill_sheet(people, [PERSON_HEADER] + [random_person(rng, i) for i in range(1, 31)])

    sales = wb.create_sheet("売上")
    rows = [["Month", "Tokyo", "Osaka", "Sendai"]]
    for m in range(1, 13):
        rows.append([datetime.date(2026, m, 1)] + [rng.randint(100, 1000) for _ in range(3)])
    fill_sheet(sales, rows)
    for row in sales.iter_rows(min_row=2, max_col=1):
        row[0].number_format = "yyyy-mm"

    wb.create_sheet("Empty")

    # Data that does not start at A1: two blank rows and two blank columns.
    offset = wb.create_sheet("Offset")
    for r, row in enumerate([["Key", "Value"], ["a", 1], ["b", 2], ["c", 3]], start=3):
        for c, value in enumerate(row, start=3):
            offset.cell(row=r, column=c, value=value)

    hidden = wb.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    fill_sheet(hidden, [["Secret", "Value"], ["x", 10], ["y", 20]])

    wide = wb.create_sheet("Wide")
    fill_sheet(wide, [[f"C{c}" for c in range(1, 101)]] +
               [[r * 100 + c for c in range(1, 101)] for r in range(1, 21)])

    # A chart sheet is a sheet, but holds no cells.
    chart = BarChart()
    chart.title = "Sales"
    chart.add_data(Reference(sales, min_col=2, min_row=1, max_col=4, max_row=13), titles_from_data=True)
    chart.set_categories(Reference(sales, min_col=1, min_row=2, max_row=13))
    wb.create_chartsheet("Chart").add_chart(chart)

    wb.save(path)


def write_merged_xlsx(path: Path, rng: random.Random):
    wb = Workbook()
    ws = wb.active
    ws.title = "report"
    center = Alignment(horizontal="center", vertical="center")

    # Row 1: a title merged across all columns. Rows 2-3: a two-level header.
    ws["A1"] = "Quarterly report"
    ws.merge_cells("A1:F1")
    ws["A2"] = "Region"
    ws.merge_cells("A2:A3")
    ws["B2"] = "Store"
    ws.merge_cells("B2:B3")
    ws["C2"] = "H1"
    ws.merge_cells("C2:D2")
    ws["E2"] = "H2"
    ws.merge_cells("E2:F2")
    for c, label in enumerate(["Q1", "Q2", "Q3", "Q4"], start=3):
        ws.cell(row=3, column=c, value=label)

    # Body: region names merged vertically over their stores.
    row = 4
    for region, stores in [("East", ["Sendai", "Tokyo", "Yokohama"]),
                           ("West", ["Osaka", "Kobe"]),
                           ("North", ["Sapporo"])]:
        ws.cell(row=row, column=1, value=region)
        if len(stores) > 1:
            ws.merge_cells(start_row=row, start_column=1, end_row=row + len(stores) - 1, end_column=1)
        for store in stores:
            ws.cell(row=row, column=2, value=store)
            for c in range(3, 7):
                ws.cell(row=row, column=c, value=rng.randint(100, 999))
            row += 1

    # A total row whose label spans two columns.
    ws.cell(row=row, column=1, value="Total")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    for c in range(3, 7):
        letter = "CDEF"[c - 3]
        ws.cell(row=row, column=c, value=f"=SUM({letter}4:{letter}{row - 1})")
    row += 2

    # A block merge, and a merge that reaches below the last stored value.
    ws.cell(row=row, column=1, value="Notes")
    ws.cell(row=row, column=2, value="This note spans a 2x5 block.")
    ws.merge_cells(start_row=row, start_column=2, end_row=row + 1, end_column=6)
    ws.cell(row=row + 3, column=1, value="Tail")
    ws.merge_cells(start_row=row + 3, start_column=1, end_row=row + 5, end_column=3)

    for cells in ws.merged_cells.ranges:
        ws.cell(row=cells.min_row, column=cells.min_col).alignment = center
    for cell in ws[2] + ws[3]:
        cell.font = Font(bold=True)
    ws["A1"].font = Font(bold=True, size=14)

    # A plain sheet with a header row merged over two columns.
    simple = wb.create_sheet("header_merge")
    simple["A1"] = "Name"
    simple["B1"] = "Contact"
    simple.merge_cells("B1:C1")
    simple.append(["Alice", "alice@example.com", "022-000-0000"])
    simple.append(["Bob", "bob@example.com", "03-0000-0000"])

    wb.save(path)


MULTILINGUAL_NAME = "多言語サンプル"

# (native name, code, script, direction, "hello", "thank you", digits 0-9)
LANGUAGES = [
    ("日本語", "ja", "Han/Hiragana/Katakana", "LTR", "こんにちは", "ありがとう", "０１２３４５６７８９"),
    ("English", "en", "Latin", "LTR", "Hello", "Thank you", "0123456789"),
    ("简体中文", "zh-Hans", "Han (Simplified)", "LTR", "你好", "谢谢", "〇一二三四五六七八九"),
    ("繁體中文", "zh-Hant", "Han (Traditional)", "LTR", "你好", "謝謝", "〇一二三四五六七八九"),
    ("한국어", "ko", "Hangul", "LTR", "안녕하세요", "감사합니다", "0123456789"),
    ("Русский", "ru", "Cyrillic", "LTR", "Здравствуйте", "Спасибо", "0123456789"),
    ("Українська", "uk", "Cyrillic", "LTR", "Привіт", "Дякую", "0123456789"),
    ("Ελληνικά", "el", "Greek", "LTR", "Γεια σας", "Ευχαριστώ", "0123456789"),
    ("العربية", "ar", "Arabic", "RTL", "مرحبا", "شكرا", "٠١٢٣٤٥٦٧٨٩"),
    ("فارسی", "fa", "Arabic (Persian)", "RTL", "سلام", "متشکرم", "۰۱۲۳۴۵۶۷۸۹"),
    ("עברית", "he", "Hebrew", "RTL", "שלום", "תודה", "0123456789"),
    ("हिन्दी", "hi", "Devanagari", "LTR", "नमस्ते", "धन्यवाद", "०१२३४५६७८९"),
    ("বাংলা", "bn", "Bengali", "LTR", "নমস্কার", "ধন্যবাদ", "০১২৩৪৫৬৭৮৯"),
    ("தமிழ்", "ta", "Tamil", "LTR", "வணக்கம்", "நன்றி", "௦௧௨௩௪௫௬௭௮௯"),
    ("ไทย", "th", "Thai", "LTR", "สวัสดี", "ขอบคุณ", "๐๑๒๓๔๕๖๗๘๙"),
    ("Tiếng Việt", "vi", "Latin (Vietnamese)", "LTR", "Xin chào", "Cảm ơn", "0123456789"),
    ("ქართული", "ka", "Georgian", "LTR", "გამარჯობა", "მადლობა", "0123456789"),
    ("Հայերեն", "hy", "Armenian", "LTR", "Բարեւ", "Շնորհակալություն", "0123456789"),
    ("አማርኛ", "am", "Ethiopic", "LTR", "ሰላም", "አመሰግናለሁ", "0123456789"),
    ("Deutsch", "de", "Latin", "LTR", "Hallo", "Danke", "0123456789"),
    ("Français", "fr", "Latin", "LTR", "Bonjour", "Merci", "0123456789"),
    ("Español", "es", "Latin", "LTR", "Hola", "Gracias", "0123456789"),
    ("Português", "pt", "Latin", "LTR", "Olá", "Obrigado", "0123456789"),
    ("Türkçe", "tr", "Latin", "LTR", "Merhaba", "Teşekkürler", "0123456789"),
    ("Polski", "pl", "Latin", "LTR", "Dzień dobry", "Dziękuję", "0123456789"),
]

# (label, text, note) -- strings that commonly break text rendering or width calculation.
UNICODE_EDGE_CASES = [
    ("emoji", "😀 🎉 🍣 🗾", "astral-plane emoji"),
    ("emoji ZWJ sequence", "👨\u200d👩\u200d👧\u200d👦", "one glyph, seven code points"),
    ("emoji flags", "🇯🇵 🇺🇸 🇩🇪 🇧🇷", "regional indicator pairs"),
    ("emoji skin tone", "👍🏻 👍🏽 👍🏿", "emoji modifiers"),
    ("combining marks", "e\u0301 a\u0300 n\u0303 / é à ñ", "decomposed (NFD) vs precomposed (NFC)"),
    ("astral-plane kanji", "𠮷野家 𩸽", "outside the BMP"),
    ("full/half width", "ＡＢＣ１２３ ｶﾀｶﾅ ABC123 カタカナ", "fullwidth Latin, halfwidth katakana"),
    ("bidirectional", "שלום world 123 مرحبا", "RTL and LTR in one cell"),
    ("zero-width", "zero\u200bwidth\u200bspace", "U+200B between words"),
    ("German sharp s", "Straße / STRASSE", "case mapping changes length"),
    ("Turkish dotted i", "İstanbul / ıiIİ", "locale-specific case mapping"),
    ("very long", "寿限無寿限無五劫の擦り切れ" * 8, "long CJK text without spaces"),
]


def make_multilingual():
    rows = [["ID", "Language", "Code", "Script", "Direction", "Hello", "ThankYou", "Digits", "Note"]]
    for i, (name, code, script, direction, hello, thanks, digits) in enumerate(LANGUAGES, start=1):
        rows.append([i, name, code, script, direction, hello, thanks, digits, ""])
    for i, (label, text, note) in enumerate(UNICODE_EDGE_CASES, start=len(LANGUAGES) + 1):
        rows.append([i, "(edge case)", "", "", "", text, "", "", f"{label}: {note}"])
    return rows


def write_multilingual_xlsx(path: Path, rows):
    """The main table, plus one sheet per language with a native sheet name."""
    wb = Workbook()
    ws = wb.active
    ws.title = MULTILINGUAL_NAME
    fill_sheet(ws, rows)
    for name, code, script, direction, hello, thanks, digits in LANGUAGES:
        if name in ("English", "日本語"):
            continue
        sheet = wb.create_sheet(name[:31])  # Excel limits sheet names to 31 characters
        fill_sheet(sheet, [["Key", "Value"], ["Language", name], ["Script", script],
                           ["Hello", hello], ["ThankYou", thanks], ["Digits", digits]])
        if direction == "RTL":
            sheet.sheet_view.rightToLeft = True
    emoji = wb.create_sheet("😀 Emoji")
    fill_sheet(emoji, [["Label", "Text"]] + [[label, text] for label, text, _ in UNICODE_EDGE_CASES])
    wb.save(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("-o", "--output", type=Path,
                        default=Path(__file__).resolve().parent / "data",
                        help="output directory (default: demo/data)")
    parser.add_argument("--no-super-large", action="store_true",
                        help="skip the super-large files (slow to generate)")
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)

    rng = random.Random(SEED)
    datasets = [
        ("normal", make_normal(rng), True),
        ("ragged", make_ragged(rng), False),
        ("large", make_large(rng), True),
    ]
    # Separate generators so that adding datasets does not change existing ones.
    datasets.append(("dates", make_dates(random.Random(SEED + 1)), True))
    datasets.append((MULTILINGUAL_NAME, make_multilingual(), True))

    for name, rows, header in datasets:
        write_csv(args.output / f"{name}.csv", to_csv_rows(rows))
        if name == "dates":
            write_dates_xlsx(args.output / f"{name}.xlsx", rows)
        elif name == MULTILINGUAL_NAME:
            write_multilingual_xlsx(args.output / f"{name}.xlsx", rows)
        else:
            write_xlsx(args.output / f"{name}.xlsx", rows, name, header=header)
        write_parquet(args.output / f"{name}.parquet", rows, header)
        print(f"wrote {name}.csv, {name}.xlsx and {name}.parquet ({len(rows)} rows)")

    normal_rows = to_csv_rows(next(rows for name, rows, _ in datasets if name == "normal"))
    for suffix in COMPRESSORS:
        write_compressed(args.output / f"normal.csv{suffix}", csv_bytes(normal_rows))
    large_rows = to_csv_rows(next(rows for name, rows, _ in datasets if name == "large"))
    write_compressed(args.output / "large.tsv.gz", csv_bytes(large_rows, "\t"))
    print("wrote normal.csv.{gz,bz2,xz,zst} and large.tsv.gz")
    write_encoded(args.output, normal_rows)
    print("wrote normal.csv in other encodings and cp932-sample.csv")
    write_json(args.output, next(rows for name, rows, _ in datasets if name == "normal"))
    print("wrote normal.json, normal.ndjson, normal.jsonl.gz and nested.json")

    dates_rows = next(rows for name, rows, _ in datasets if name == "dates")
    write_dates_xlsx(args.output / "dates1904.xlsx", dates_rows, epoch=CALENDAR_MAC_1904)
    print("wrote dates1904.xlsx")
    write_multisheet_xlsx(args.output / "multisheet.xlsx", random.Random(SEED + 2))
    print("wrote multisheet.xlsx")
    write_merged_xlsx(args.output / "merged.xlsx", random.Random(SEED + 3))
    print("wrote merged.xlsx")

    tables = ", ".join(name for name, _, _ in datasets)
    for db_name, writer in [("demo.sqlite3", write_sqlite), ("demo.duckdb", write_duckdb)]:
        writer(args.output / db_name, datasets)
        print(f"wrote {db_name} ({tables})")

    if not args.no_super_large:
        write_super_large(args.output, SUPER_LARGE_SEED)
        print(f"wrote super-large.csv.xz, .xlsx and .parquet ({SUPER_LARGE_ROWS} data rows)")


if __name__ == "__main__":
    main()
